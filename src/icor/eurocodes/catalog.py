"""Parse ICOR's Eurocode 4 catalogue into deterministic, auditable rows.

The workbook is free text maintained by hand: brands carry trailing spaces and
Dutch qualifiers, one model cell can describe several vehicles and their start
years (`ALPHA 5D HBK 2008/BETA 5D SUV 2010-`), and the end-of-production column
is empty or zero for most rows. Nothing here guesses past what a row says. Each
row is either parsed, with every year, drive side, body word and market note it
states, or rejected with a reason, so parsed plus rejected always equals the
rows read. Matching rows to vehicles happens later, under review.

The runtime never parses the workbook: an ingest script writes the parsed
catalogue to a private JSON document with the source's sha256.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from icor.evidence.vehicle_identity import fold_make

HEADER = ("Eurocode 4", "Brand", "Model", "Eop Year", "Sop Year")
DOCUMENT_FORMAT = "icor-eurocode-catalogue"
DOCUMENT_VERSION = 1

FIRST_YEAR = 1900
LAST_YEAR = 2039

# Brand qualifiers that state a scope rather than a make.
_BRAND_SCOPE_WORDS = {
    "vrachtwagen": "truck",
    "vrachtwagens": "truck",
    "trucks": "truck",
    "bus": "bus",
    "usa": "non_european",
    "us": "non_european",
}
# Brand qualifiers that only say the brand covers cars and vans.
_BRAND_NOISE = re.compile(r"\bauto'?s\s*&\s*vans\b", re.IGNORECASE)

_BODY_WORDS = {
    "HBK": "hatchback",
    "HATCHBACK": "hatchback",
    "SAL": "saloon",
    "SALOON": "saloon",
    "SEDAN": "saloon",
    "SED": "saloon",
    "LIMOUSINE": "saloon",
    "BRK": "estate",
    "BREAK": "estate",
    "EST": "estate",
    "ESTATE": "estate",
    "VARIANT": "estate",
    "AVANT": "estate",
    "TOURING": "estate",
    "COMBI": "estate",
    "KOMBI": "estate",
    "SW": "estate",
    "S.W.": "estate",
    "WAGON": "estate",
    "GRANDTOUR": "estate",
    "SPORTSWAGON": "estate",
    "CPE": "coupe",
    "COUPE": "coupe",
    "CAB": "convertible",
    "CABRIO": "convertible",
    "CABRIOLET": "convertible",
    "CONVERTIBLE": "convertible",
    "ROADSTER": "convertible",
    "SPIDER": "convertible",
    "SPYDER": "convertible",
    "TARGA": "convertible",
    "SOFTTOP": "convertible",
    "SUV": "suv",
    "CROSSOVER": "suv",
    "MPV": "mpv",
    "MONOSPACE": "mpv",
    "MINIVAN": "mpv",
    "VAN": "van",
    "PU": "pickup",
    "PICK-UP": "pickup",
    "PICKUP": "pickup",
    "MINIBUS": "bus",
    "BUS": "bus",
    "TRUCK": "truck",
    "VRACHTW": "truck",
    "VRACHTWAGEN": "truck",
}
# `CAB` after these words is a pickup or truck cabin, not a convertible.
_CABIN_PREFIXES = frozenset(
    {"DOUBLE", "SINGLE", "SUPER", "CREW", "KING", "REGULAR", "CHASSIS", "DAY", "SLEEPER"}
)
_DRIVE_WORDS = {"LHD": "LHD", "LINKS": "LHD", "RHD": "RHD", "RECHTS": "RHD"}
# Words naming a market or origin. Followed by a production word they say where
# the car was built (`JAPANESE PRODUCTION`), which can still be a European car;
# otherwise they name a version sold elsewhere (`JAPANESE MARKET`, `(BRAZIL)`).
_ORIGIN_WORDS = frozenset(
    {
        "JAPANESE",
        "JAPAN",
        "JAP",
        "JAPAANSE",
        "RUSSIAN",
        "RUS",
        "BRAZIL",
        "KOREAN",
        "CHINA",
        "CHINESE",
        "INDIA",
        "INDIAN",
        "AUSTRALIA",
        "AUSTRALIAN",
        "AUSTRALIE",
        "MEXICO",
    }
)
_PRODUCTION_WORDS = frozenset({"PRODUCTION", "PRODUKTIE", "PRODUCTIE", "BUILT", "MADE"})
_NON_EUROPEAN_WORDS = frozenset({"USA", "US", "JDM", "IMPORT"})
_NON_EUROPEAN_PHRASES = ("NORTH AMERICAN", "NORTH AMERICA")
_EUROPEAN_MARKERS = re.compile(r"\(\s*EU\s*\)|\bEUROPEAN\b|\bEUROPE\b")
_CONVERSION_WORDS = frozenset({"HEARSE", "CAMPER", "AMBULANCE", "CONVERSION", "CONVERTED"})

_TOKEN = re.compile(r"[A-Z0-9][A-Z0-9.'&-]*")
_YEAR = r"(?:19\d\d|20[0-3]\d)"
_FOUR_DIGIT = re.compile(rf"(?<![\d.])({_YEAR})(?![\d.])")
# A day or month before a four-digit year: `07/2009`, `27/07/2017`, `05-2022`.
_DATE_PREFIX = re.compile(
    rf"(?<![\w./])(?:(?:0?[1-9]|[12]\d|3[01])/)?(0?[1-9]|1[0-2])([/-])(?={_YEAR}(?![\d.]))"
)
# A month before a two-digit year that opens a range: `04/99-`.
_MONTH_TWO_DIGIT = re.compile(r"(?<![\w./])(?:0?[1-9]|1[0-2])/(?=\d{2}\s*-)")
# A two-digit year opening a range: `97-`, `74-83`, `95-2006`, `97-/`. The dash
# must lead to a separator or another year, so `17-26M` and `87-E-2000` (model
# and chassis numbers) are not read.
_TWO_DIGIT_START = re.compile(
    rf"(?<![\w./])(\d{{2}})(?=\s*-\s*(?:$|[\s/,+)]|\d{{2}}(?![\w.])|{_YEAR}(?![\d.])))"
)
_TWO_DIGIT_END = re.compile(r"\A\s*-\s*(\d{2})(?![\w./])")
_NUMBER_BEFORE = re.compile(r"(?<![\w.])(\d+)\s*[-,]\s*\Z")
_NUMBER_AFTER = re.compile(r"\A\s*[-,]\s*(\d+)(?![\w.])")


@dataclass(frozen=True, slots=True)
class YearMention:
    """One year a description states, where it states it, and what it means.

    ``role`` is ``start`` for `2012-`, `2007+` or a bare `2008`, ``end`` for the
    second half of `2007-2013` or for `TOT 10/2009` (Dutch: until), and a range
    contributes one start and one end mention.
    """

    position: int
    year: int
    role: str


@dataclass(frozen=True, slots=True)
class CatalogueRow:
    row: int
    code: str
    brand: str
    makes: tuple[str, ...]
    description: str
    sop: int | None
    eop: int | None
    start: int | None
    end: int | None
    drive: str | None
    bodies: tuple[str, ...]
    scope: str
    years: tuple[YearMention, ...]
    flags: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class RejectedRow:
    row: int
    reason: str
    values: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class Catalogue:
    rows: tuple[CatalogueRow, ...]
    rejected: tuple[RejectedRow, ...]
    source_sha256: str
    source_rows: int

    def codes(self) -> dict[str, tuple[CatalogueRow, ...]]:
        grouped: dict[str, list[CatalogueRow]] = {}
        for row in self.rows:
            grouped.setdefault(row.code, []).append(row)
        return {code: tuple(rows) for code, rows in sorted(grouped.items())}


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return " ".join(text.split())


def _year_or_none(value: object) -> tuple[int | None, str | None]:
    """Read a year cell. Zero and empty mean unknown; anything else must be a year."""

    if value is None:
        return None, None
    text = _clean_text(value)
    if text in {"", "0"}:
        return None, None
    try:
        number = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None, "unreadable_year"
    if number == 0:
        return None, None
    if not FIRST_YEAR <= number <= LAST_YEAR:
        return None, "year_out_of_range"
    return number, None


def parse_brand(raw: str, make_aliases: Mapping[str, str]) -> tuple[tuple[str, ...], str | None]:
    """Fold a catalogue brand to app makes and the scope its qualifier states.

    `ALPHA/BETA` names two rebadged makes and yields both. A qualifier that only
    says the brand also covers vans is dropped; one that says trucks, buses or a
    North American line becomes the row's scope instead of part of the make.
    """

    text = _BRAND_NOISE.sub(" ", _clean_text(raw)).casefold()
    makes: list[str] = []
    scope: str | None = None
    for part in text.split("/"):
        words = part.split()
        while len(words) > 1 and words[-1] in _BRAND_SCOPE_WORDS:
            scope = _BRAND_SCOPE_WORDS[words.pop()]
        name = " ".join(words)
        if not name:
            continue
        make = make_aliases.get(name) or fold_make(name)
        if make and make not in makes:
            makes.append(make)
    return tuple(makes), scope


def _is_model_number(text: str, start: int, end: int) -> bool:
    """A four-digit number that opens a description or a `/` segment names a model.

    `9193 5D SUV 2021-` describes a vehicle called 9193; `HBK 2008/BETA` is a
    start year. The number is a model only when it begins the text or follows a
    slash and is followed by a word rather than a dash, a plus or the end.
    """

    before = text[:start].rstrip()
    if before and not before.endswith("/"):
        return False
    if re.search(r"(?<!\d)\d{1,2}/$", before):
        # `04/2007` and `27/07/2017` are dates: the slash follows a month or day.
        return False
    after = text[end:]
    if not after.strip() or after.lstrip().startswith(("-", "+", "/")):
        return False
    return after.startswith(" ")


def _is_year_like(digits: str) -> bool:
    return len(digits) == 2 or bool(re.fullmatch(_YEAR, digits))


def _in_number_chain(text: str, start: int, end: int) -> bool:
    """`1300-2000`, `2000-2600` and `2000,2500` are engine or model numbers.

    A four-digit year joined by a dash or a comma to a number that cannot be a
    year is part of that list, not a production date.
    """

    before = _NUMBER_BEFORE.search(text[:start])
    after = _NUMBER_AFTER.search(text[end:])
    return bool(
        (before and not _is_year_like(before.group(1)))
        or (after and not _is_year_like(after.group(1)))
    )


def _blank_date_prefixes(text: str, sop: int | None) -> str:
    """Replace day and month prefixes with spaces, keeping every position.

    `05-2022` is May 2022, unless the start column says the row starts in 2005,
    in which case it is the range `2005-2022`.
    """

    def blank(match: re.Match[str]) -> str:
        if match.group(2) == "-" and sop == 2000 + int(match.group(1)):
            return match.group(0)
        return " " * len(match.group(0))

    return _MONTH_TWO_DIGIT.sub(
        lambda match: " " * len(match.group(0)), _DATE_PREFIX.sub(blank, text)
    )


def _years(text: str, sop: int | None = None) -> tuple[tuple[YearMention, ...], tuple[str, ...]]:
    scan = _blank_date_prefixes(text, sop)
    flags: list[str] = []
    found: list[tuple[int, int, int, bool]] = []
    for match in _FOUR_DIGIT.finditer(scan):
        glued = match.start() and scan[match.start() - 1].isalpha()
        if glued and not scan[match.end() :].lstrip().startswith(("-", "+")):
            # `E2000 VAN`, `L2000`: a model name. `BRK2007-` is a glued start year.
            flags.append("glued_number_read_as_model")
            continue
        if _is_model_number(scan, match.start(), match.end()):
            flags.append("leading_number_read_as_model")
            continue
        if _in_number_chain(scan, match.start(), match.end()):
            flags.append("number_list_read_as_model")
            continue
        found.append((match.start(), match.end(), int(match.group(1)), False))
    for match in _TWO_DIGIT_START.finditer(scan):
        found.append((match.start(1), match.end(1), _expand_two_digit(match.group(1)), False))
        closing = _TWO_DIGIT_END.match(scan[match.end(1) :])
        if closing:
            position = match.end(1) + closing.start(1)
            found.append((position, position + 2, _expand_two_digit(closing.group(1)), True))
    found.sort()
    mentions: list[YearMention] = []
    previous_end: int | None = None
    for position, end, year, closes in found:
        between = scan[previous_end:position] if previous_end is not None else None
        until = re.search(r"\b(?:TOT|UNTIL)\s*\Z", scan[:position]) is not None
        closes_range = (
            between is not None
            and re.fullmatch(r"\s*-\s*", between) is not None
            and mentions[-1].role == "start"
        )
        role = "end" if closes or until or closes_range else "start"
        mentions.append(YearMention(position, year, role))
        previous_end = end
    return tuple(mentions), tuple(dict.fromkeys(flags))


def _expand_two_digit(value: str) -> int:
    """`97-` is 1997 and `05-` is 2005; the catalogue writes recent years in full."""

    number = int(value)
    return 2000 + number if number < 30 else 1900 + number


def _tokens(text: str) -> list[str]:
    return _TOKEN.findall(text)


_DOORS = re.compile(r"\d(?:D|P)\Z")


def _bodies(tokens: Sequence[str]) -> tuple[tuple[str, ...], tuple[str, ...]]:
    found: list[str] = []
    flags: list[str] = []
    for position, token in enumerate(tokens):
        body = _BODY_WORDS.get(token)
        if body is None:
            continue
        previous = tokens[position - 1] if position else ""
        following = tokens[position + 1] if position + 1 < len(tokens) else ""
        if token == "CAB" and previous in _CABIN_PREFIXES:
            body = "pickup"
        if token == "VAN" and not _DOORS.match(previous):
            # Dutch `van` means "from" (`VAN DAM 1281081`, `HBK VAN 10/2009-`),
            # so only a door count (`3D VAN`) or a non-date word settles it.
            if following in {"DAM", "CHASSIS", "CH"}:
                continue
            if following[:1].isdigit():
                flags.append("van_word_ambiguous")
                continue
        if body not in found:
            found.append(body)
    joined = " ".join(tokens)
    if ("SPORTS TOURER" in joined or "GRAND TOUR" in joined) and "estate" not in found:
        found.append("estate")
    return tuple(found), tuple(flags)


def _drive(tokens: Sequence[str]) -> tuple[str | None, tuple[str, ...]]:
    sides = {_DRIVE_WORDS[token] for token in tokens if token in _DRIVE_WORDS}
    if len(sides) > 1:
        return None, ("both_drive_sides",)
    return (next(iter(sides)) if sides else None), ()


def _origin(tokens: Sequence[str], text: str) -> tuple[bool, bool]:
    """Whether the row names a non-European version, and whether it names a
    production origin only."""

    non_european = any(word in tokens for word in _NON_EUROPEAN_WORDS) or any(
        phrase in text for phrase in _NON_EUROPEAN_PHRASES
    )
    production_origin = False
    for position, token in enumerate(tokens):
        if token not in _ORIGIN_WORDS:
            continue
        following = tokens[position + 1] if position + 1 < len(tokens) else ""
        if following in _PRODUCTION_WORDS:
            production_origin = True
        else:
            non_european = True
    if _EUROPEAN_MARKERS.search(text):
        non_european = False
    return non_european, production_origin


def _row_scope(brand_scope: str | None, tokens: Sequence[str], non_european: bool) -> str:
    if brand_scope is not None:
        return brand_scope
    if "VRACHTW" in tokens or "VRACHTWAGEN" in tokens:
        return "truck"
    if any(word in tokens for word in _CONVERSION_WORDS):
        return "conversion"
    if non_european:
        return "non_european"
    return "passenger"


def parse_rows(
    rows: Iterable[Sequence[object]],
    *,
    make_aliases: Mapping[str, str],
    source_sha256: str,
    first_row: int = 2,
) -> Catalogue:
    """Parse data rows (header excluded); ``first_row`` is the sheet row of the first."""

    parsed: list[CatalogueRow] = []
    rejected: list[RejectedRow] = []
    count = 0
    for offset, values in enumerate(rows):
        count += 1
        row_number = first_row + offset
        cells = tuple(values) + (None,) * (5 - len(tuple(values)))
        raw_code, raw_brand, raw_model, raw_eop, raw_sop = cells[:5]
        texts = tuple(_clean_text(cell) for cell in cells[:5])
        code = _clean_text(raw_code).upper()
        if not code and not any(texts[1:3]) and texts[3] in {"", "0"} and texts[4] in {"", "0"}:
            rejected.append(RejectedRow(row_number, "blank_row", texts))
            continue
        if not code:
            rejected.append(RejectedRow(row_number, "missing_code", texts))
            continue
        if len(code) != 4 or not code.isalnum():
            rejected.append(RejectedRow(row_number, "code_not_four_characters", texts))
            continue
        brand = _clean_text(raw_brand)
        if not brand:
            rejected.append(RejectedRow(row_number, "missing_brand", texts))
            continue
        description = _clean_text(raw_model).upper()
        if not description:
            rejected.append(RejectedRow(row_number, "missing_model", texts))
            continue
        makes, brand_scope = parse_brand(brand, make_aliases)
        if not makes:
            rejected.append(RejectedRow(row_number, "unreadable_brand", texts))
            continue

        flags: list[str] = []
        sop, sop_problem = _year_or_none(raw_sop)
        eop, eop_problem = _year_or_none(raw_eop)
        if sop_problem:
            flags.append(f"sop_{sop_problem}")
        if eop_problem:
            flags.append(f"eop_{eop_problem}")
        if sop is None and not sop_problem:
            flags.append("sop_missing")
        if sop is not None and eop is not None and eop < sop:
            flags.append("eop_before_sop")
            eop = None

        years, year_flags = _years(description, sop)
        flags.extend(year_flags)
        stated_end = (
            eop
            if eop is not None
            else max((mention.year for mention in years if mention.role == "end"), default=None)
        )
        if stated_end is not None and any(
            mention.role == "start" and mention.year > stated_end for mention in years
        ):
            # `ZEPHYRA 1990 SAL 79-85`: a start after the row's own end is a
            # model or engine number, never a production year.
            years = tuple(
                mention
                for mention in years
                if not (mention.role == "start" and mention.year > stated_end)
            )
            flags.append("start_after_end_dropped")
        starts = [mention.year for mention in years if mention.role == "start"]
        ends = [mention.year for mention in years if mention.role == "end"]
        if sop is not None and starts and sop not in starts:
            flags.append("sop_not_in_text")
        start_candidates = ([sop] if sop is not None else []) + starts
        start = min(start_candidates) if start_candidates else None
        end = eop
        if end is None and ends and len(ends) >= len(starts):
            end = max(ends)
        if start is not None and end is not None and end < start:
            flags.append("end_before_start")
            end = None
        if start is None:
            flags.append("no_start_year")

        tokens = _tokens(description)
        bodies, body_flags = _bodies(tokens)
        flags.extend(body_flags)
        drive, drive_flags = _drive(tokens)
        flags.extend(drive_flags)
        non_european, production_origin = _origin(tokens, description)
        if production_origin:
            flags.append("production_origin")
        scope = _row_scope(brand_scope, tokens, non_european)
        parsed.append(
            CatalogueRow(
                row=row_number,
                code=code,
                brand=brand,
                makes=makes,
                description=description,
                sop=sop,
                eop=eop,
                start=start,
                end=end,
                drive=drive,
                bodies=bodies,
                scope=scope,
                years=years,
                flags=tuple(dict.fromkeys(flags)),
            )
        )
    return Catalogue(tuple(parsed), tuple(rejected), source_sha256, count)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def read_workbook(path: Path, *, make_aliases: Mapping[str, str]) -> Catalogue:
    """Read the catalogue workbook, refusing any sheet whose header differs."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("the Eurocode catalogue must have exactly one sheet")
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header = tuple(_clean_text(cell) for cell in next(iterator, ()))
        if header[: len(HEADER)] != HEADER or any(header[len(HEADER) :]):
            raise ValueError("unexpected Eurocode catalogue header")
        rows = [tuple(values[: len(HEADER)]) for values in iterator]
    finally:
        workbook.close()
    return parse_rows(rows, make_aliases=make_aliases, source_sha256=file_sha256(path))


def to_document(catalogue: Catalogue) -> dict[str, Any]:
    """A deterministic JSON-ready document; same catalogue, same bytes."""

    return {
        "format": DOCUMENT_FORMAT,
        "version": DOCUMENT_VERSION,
        "source": {"sha256": catalogue.source_sha256, "rows": catalogue.source_rows},
        "rows": [
            {
                "row": row.row,
                "code": row.code,
                "brand": row.brand,
                "makes": list(row.makes),
                "description": row.description,
                "sop": row.sop,
                "eop": row.eop,
                "start": row.start,
                "end": row.end,
                "drive": row.drive,
                "bodies": list(row.bodies),
                "scope": row.scope,
                "years": [[mention.position, mention.year, mention.role] for mention in row.years],
                "flags": list(row.flags),
            }
            for row in sorted(catalogue.rows, key=lambda row: row.row)
        ],
        "rejected": [
            {"row": row.row, "reason": row.reason, "values": list(row.values)}
            for row in sorted(catalogue.rejected, key=lambda row: row.row)
        ],
    }


def from_document(document: Mapping[str, Any]) -> Catalogue:
    if document.get("format") != DOCUMENT_FORMAT or document.get("version") != DOCUMENT_VERSION:
        raise ValueError("not an ICOR Eurocode catalogue document")
    rows = tuple(
        CatalogueRow(
            row=int(item["row"]),
            code=str(item["code"]),
            brand=str(item["brand"]),
            makes=tuple(item["makes"]),
            description=str(item["description"]),
            sop=item["sop"],
            eop=item["eop"],
            start=item["start"],
            end=item["end"],
            drive=item["drive"],
            bodies=tuple(item["bodies"]),
            scope=str(item["scope"]),
            years=tuple(YearMention(int(p), int(y), str(r)) for p, y, r in item["years"]),
            flags=tuple(item["flags"]),
        )
        for item in document["rows"]
    )
    rejected = tuple(
        RejectedRow(int(item["row"]), str(item["reason"]), tuple(item["values"]))
        for item in document["rejected"]
    )
    source = document["source"]
    catalogue = Catalogue(rows, rejected, str(source["sha256"]), int(source["rows"]))
    if len(rows) + len(rejected) != catalogue.source_rows:
        raise ValueError("catalogue document does not account for every source row")
    return catalogue
